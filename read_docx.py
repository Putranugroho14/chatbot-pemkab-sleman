from docx import Document
import json

doc = Document(r'd:\SEMESTER 6\Chatbot\laporan magang.docx')

output = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        output.append({
            "index": i,
            "style": para.style.name,
            "text": text[:300]
        })

# Print all paragraphs
for item in output:
    print(f"[{item['index']}] STYLE={item['style']}")
    print(f"     TEXT={item['text']}")
    print()

print(f"\n=== TOTAL paragraphs with text: {len(output)} ===")

# Also check tables
print(f"\n=== TABLES: {len(doc.tables)} ===")
for ti, table in enumerate(doc.tables):
    print(f"Table {ti}: {len(table.rows)} rows x {len(table.columns)} cols")
    for ri, row in enumerate(table.rows):
        if ri < 3:
            cells = [c.text.strip()[:50] for c in row.cells]
            print(f"  Row {ri}: {cells}")
