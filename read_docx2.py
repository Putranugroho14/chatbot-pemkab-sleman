from docx import Document

doc = Document(r'd:\SEMESTER 6\Chatbot\laporan magang.docx')

with open(r'd:\SEMESTER 6\Chatbot\laporan_content.txt', 'w', encoding='utf-8') as f:
    f.write("=== PARAGRAPHS ===\n\n")
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            f.write(f"[{i}] STYLE={para.style.name}\n")
            f.write(f"     TEXT={text[:400]}\n\n")
    
    f.write(f"\n=== TOTAL paragraphs: {sum(1 for p in doc.paragraphs if p.text.strip())} ===\n\n")
    
    f.write("=== TABLES ===\n\n")
    for ti, table in enumerate(doc.tables):
        f.write(f"Table {ti}: {len(table.rows)} rows x {len(table.columns)} cols\n")
        for ri, row in enumerate(table.rows):
            cells = [c.text.strip()[:80] for c in row.cells]
            f.write(f"  Row {ri}: {cells}\n")
        f.write("\n")

print("Done! Content written to laporan_content.txt")
